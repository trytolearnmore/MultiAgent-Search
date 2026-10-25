from pathlib import Path
import sys
from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

path = Path(r"C:\Users\86183\Desktop\项目经历.docx")
doc = Document(path)
print("PARAGRAPHS")
for i, p in enumerate(doc.paragraphs, 1):
    text = p.text.strip()
    if text:
        print(f"P{i:03d} [{p.style.name}] {text}")
print("\nTABLES")
for ti, table in enumerate(doc.tables, 1):
    print(f"TABLE {ti}")
    for ri, row in enumerate(table.rows, 1):
        vals = []
        for cell in row.cells:
            vals.append(" / ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()))
        print(f"R{ri:02d}: " + " || ".join(vals))
print("\nTEXTBOX_OR_BODY_XML_TEXT")
body = doc.element.body
for i, node in enumerate(body.iter(), 1):
    if node.tag == qn('w:t') and node.text and node.text.strip():
        print(node.text.strip())
